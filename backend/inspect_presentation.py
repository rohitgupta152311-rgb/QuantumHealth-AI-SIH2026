import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = pptx.Presentation(r'C:\Users\rohit\OneDrive\Desktop\SIH_2026_QuantumHealth_AI_NIT_Nagaland.pptx')

with open(r'backend\slides_detailed_map.txt', 'w', encoding='utf-8') as f:
    for s_idx, slide in enumerate(prs.slides):
        f.write(f"\n{'='*30} SLIDE {s_idx + 1} (total shapes: {len(slide.shapes)}) {'='*30}\n")
        for sh_idx, shape in enumerate(slide.shapes):
            f.write(f"\n--- Shape [{sh_idx}] id={shape.shape_id}, name='{shape.name}', type={shape.shape_type} ---\n")
            f.write(f"    Left: {shape.left}, Top: {shape.top}, Width: {shape.width}, Height: {shape.height}\n")
            if shape.has_text_frame:
                tf = shape.text_frame
                f.write(f"    Paragraphs count: {len(tf.paragraphs)}\n")
                for p_idx, p in enumerate(tf.paragraphs):
                    runs_info = " | ".join([f"'{r.text}'(sz={r.font.size.pt if r.font.size else 'None'}pt, b={r.font.bold})" for r in p.runs])
                    f.write(f"      [P{p_idx}] text='{p.text}'\n")
                    f.write(f"             runs: {runs_info}\n")
            elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                f.write("    (Group shape)\n")
            elif shape.has_table:
                f.write("    (Table shape)\n")
                for r_idx, row in enumerate(shape.table.rows):
                    row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
                    f.write(f"      Row {r_idx}: {row_txt}\n")

print("Detailed map written to backend/slides_detailed_map.txt successfully.")
