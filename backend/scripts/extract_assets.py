import sys
import os
import pptx

prs = pptx.Presentation(r'C:\Users\rohit\OneDrive\Attachments\Desktop\SIH_2026_QuantumHealth_AI_NIT_Nagaland.pptx')
os.makedirs('backend/scripts/slide_assets', exist_ok=True)
pic_count = 0
for s_idx, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            pic_count += 1
            image = shape.image
            ext = image.ext
            clean_name = shape.name.replace(' ', '_').replace(':', '_')
            filename = f'backend/scripts/slide_assets/slide_{s_idx}_{clean_name}.{ext}'
            with open(filename, 'wb') as f:
                f.write(image.blob)
            print(f'Saved {filename} ({len(image.blob)} bytes) from slide {s_idx}: {shape.left}, {shape.top}, {shape.width}, {shape.height}')
print(f'Total pictures extracted: {pic_count}')
